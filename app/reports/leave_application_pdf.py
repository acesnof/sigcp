"""Geração do Application Form For Leave a partir do modelo oficial."""

from datetime import datetime
from pathlib import Path
import shutil
import subprocess
import tempfile

from docx import Document
from docx.shared import Mm, Pt

from app.config import DOCS_DIR


TEMPLATE_NAME = "application_form_for_leave.docx"


def _date(value):
    try:
        return datetime.strptime(str(value)[:10], "%Y-%m-%d").strftime("%d/%m/%Y")
    except ValueError:
        return ""


def _set_paragraph_text(paragraph, value):
    if not paragraph.runs:
        return
    paragraph.runs[0].text = str(value)
    for run in paragraph.runs[1:]:
        run.text = ""


def _set_labeled_value(paragraph, label, value):
    original_run = next((run for run in paragraph.runs if run.text.strip()), None)
    original_size = original_run.font.size if original_run else None
    paragraph.clear()
    label_run = paragraph.add_run(label)
    if original_size:
        label_run.font.size = original_size
    value_run = paragraph.add_run(f" {value}")
    value_run.bold = False
    if original_size:
        value_run.font.size = original_size


def _set_underlined_labeled_value(paragraph, label, value, line):
    original_run = next((run for run in paragraph.runs if run.text.strip()), None)
    original_size = original_run.font.size if original_run else None
    paragraph.clear()
    label_run = paragraph.add_run(label)
    label_run.bold = True
    if original_size:
        label_run.font.size = original_size
    if value:
        value_run = paragraph.add_run(f" {value}")
        value_run.bold = False
        value_run.underline = True
        if original_size:
            value_run.font.size = original_size
    else:
        line_run = paragraph.add_run(line)
        if original_size:
            line_run.font.size = original_size


def _set_labeled_value_with_line(paragraph, label, value, line):
    original_run = next((run for run in paragraph.runs if run.text.strip()), None)
    original_size = original_run.font.size if original_run else None
    paragraph.clear()
    label_run = paragraph.add_run(label)
    if original_size:
        label_run.font.size = original_size
    if value:
        value_run = paragraph.add_run(f" {value}")
        value_run.bold = False
        value_run.underline = True
        if original_size:
            value_run.font.size = original_size
    line_run = paragraph.add_run(line)
    if original_size:
        line_run.font.size = original_size


def _set_snr_signature(paragraph, value, flag_path, digital_validation):
    original_run = next((run for run in paragraph.runs if run.text.strip()), None)
    label_size = original_run.font.size if original_run else None
    paragraph.clear()
    label = paragraph.add_run("12. COORDINATION 2: Senior National Representative:")
    label.bold = True
    if label_size:
        label.font.size = label_size
    if digital_validation:
        space = paragraph.add_run(" ")
        flag = paragraph.add_run()
        flag.add_picture(str(flag_path), width=Mm(4.5))
        signature = paragraph.add_run(f" {value}")
        signature.bold = False
        signature.font.size = max(label_size - Pt(3), Pt(4)) if label_size else Pt(7)
    if not digital_validation:
        line = paragraph.add_run()
        line.add_text(" ________________________")
        if label_size:
            line.font.size = label_size


def _paragraphs_in_table(table):
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested in cell.tables:
                yield from _paragraphs_in_table(nested)


def _replace_mission_name(document, mission_name):
    mission_name = str(mission_name or "").strip().upper()
    if not mission_name:
        return
    for section in document.sections:
        for footer in (section.footer, section.first_page_footer, section.even_page_footer):
            paragraphs = list(footer.paragraphs)
            for table in footer.tables:
                paragraphs.extend(_paragraphs_in_table(table))
            for paragraph in paragraphs:
                for run in paragraph.runs:
                    if "EUTM RCA" in run.text:
                        run.text = run.text.replace("EUTM RCA", mission_name)


def _replace_template_values(docx_path, values):
    document = Document(docx_path)
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        paragraphs.extend(_paragraphs_in_table(table))
    for paragraph in paragraphs:
        current = paragraph.text
        if current in values:
            _set_paragraph_text(paragraph, values[current])
        elif current.startswith("8. LEAVE ADDRESS:"):
            _set_underlined_labeled_value(
                paragraph, "8. LEAVE ADDRESS:", values["leave_address"], " _______________________"
            )
        elif current.startswith("9. Telephone number of contact:"):
            _set_underlined_labeled_value(
                paragraph, "9. Telephone number of contact:", values["phone"], " _______________________"
            )
        elif current.startswith("POC NOMINATED BY CHIEF DURING ABSENCE:"):
            _set_underlined_labeled_value(
                paragraph,
                "POC NOMINATED BY CHIEF DURING ABSENCE:",
                values["poc_nominated"],
                " _______________________",
            )
        elif current.startswith("SIGNATURE OF POC TAKING KNOWLODGE:"):
            _set_underlined_labeled_value(
                paragraph,
                "SIGNATURE OF POC TAKING KNOWLODGE:",
                "",
                " _______________________________________",
            )
        elif current.startswith("REASON FOR DEVIATION FROM LEAVE PLAN (If any):"):
            _set_underlined_labeled_value(
                paragraph,
                "REASON FOR DEVIATION FROM LEAVE PLAN (If any):",
                values["deviation_reason"] or "None",
                " _______________________",
            )
        elif current.startswith("12. COORDINATION 2:"):
            _set_snr_signature(
                paragraph,
                values["snr_signature"],
                values["flag_path"],
                values["digital_validation"],
            )
    _replace_mission_name(document, values.get("mission_name"))
    document.save(docx_path)


def _convert_to_pdf(docx_path, pdf_path):
    def ps_literal(value):
        return "'" + str(value).replace("'", "''") + "'"

    script = (
        "$ErrorActionPreference='Stop'; "
        f"$inputPath={ps_literal(docx_path)}; $outputPath={ps_literal(pdf_path)}; "
        "$word=New-Object -ComObject Word.Application; $word.Visible=$false; $word.DisplayAlerts=0; "
        "try { $document=$word.Documents.OpenNoRepairDialog($inputPath); "
        "$document.ExportAsFixedFormat($outputPath,17); $document.Close($false); } "
        "finally { $word.Quit(); }"
    )
    result = subprocess.run(
        ["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", script],
        capture_output=True, text=True, timeout=90,
    )
    if result.returncode != 0 or not pdf_path.is_file():
        detail = (result.stderr or result.stdout or "Microsoft Word não devolveu o PDF.").strip()
        raise RuntimeError(f"Não foi possível gerar o PDF do pedido de férias: {detail}")


def generate_leave_application_pdf(
    destination, person, vacation, snr=None, poc_nominated="", deviation_reason="", mission_name="",
    *, digital_validation=None,
):
    """Preenche o modelo oficial e converte-o para PDF pelo Microsoft Word."""
    template = Path(DOCS_DIR) / TEMPLATE_NAME
    if not template.is_file():
        raise RuntimeError("O modelo Application Form For Leave não está disponível.")
    phone = "".join(char for char in str(person.get("telefone_contacto_pt") or "") if char.isdigit())
    phone = f"+351 {phone}" if phone else ""
    summary = vacation.get("resumo") or {}
    signature = "12. COORDINATION 2: Senior National Representative: _________________________________"
    approval_date = _date(vacation.get("aprovacao_em") or vacation.get("decidido_em"))
    if digital_validation is None:
        digital_validation = bool(snr and int(snr.get("validacao_digital_leave") or 0))
    digital_validation = bool(digital_validation and snr)
    if digital_validation:
        rank = (snr.get("posto") or "").strip()
        surname = (snr.get("sobrenome") or "").strip().upper()
        name = (snr.get("nome") or "").strip()
        signer = " ".join(part for part in (rank, f"{surname}," if surname else "", name) if part).replace(" ,", ",")
        signature = f"Digitally validated by {signer}"
        signature += f" on {approval_date}" if approval_date else " (approval date not recorded)"
    values = {
        "LDG": person.get("branch_pillar") or "",
        "Madeira Roberto": " ".join(
            part for part in ((person.get("nome") or "").strip(), (person.get("sobrenome") or "").strip()) if part
        ).upper(),
        "OR8": person.get("posto") or "",
        "PRT - ARMY": {
            "Exército": "PRT - ARMY",
            "Força Aérea": "PRT - AIR FORCE",
            "Marinha": "PRT - NAVY",
        }.get(person.get("ramo"), "PRT - ARMY"),
        "04/11/2026": _date(vacation.get("data_hora_inicio")),
        "24/11/2026": _date(vacation.get("data_hora_fim")), "100/20": person.get("posicao_numero") or "",
        "21": str(summary.get("dias_ausencia") or 0), "11": str(summary.get("dias_ferias") or 0),
        "leave_address": person.get("morada_pt") or "", "phone": phone, "snr_signature": signature,
        "poc_nominated": str(poc_nominated or "").strip(),
        "deviation_reason": str(deviation_reason or "").strip(),
        "mission_name": mission_name,
        "digital_validation": digital_validation, "flag_path": Path(DOCS_DIR) / "prt.png",
    }
    with tempfile.TemporaryDirectory(prefix="sigcp_leave_") as temporary:
        working_docx = Path(temporary) / TEMPLATE_NAME
        shutil.copy2(template, working_docx)
        _replace_template_values(working_docx, values)
        _convert_to_pdf(working_docx, Path(destination))
