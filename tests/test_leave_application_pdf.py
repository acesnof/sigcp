from pathlib import Path
from unittest.mock import patch

import pytest
from docx import Document

from app.reports import leave_application_pdf as report


@pytest.mark.parametrize("enabled,approval_date,expected", [
    (1, "2026-05-20 14:30", "20/05/2026"),
    (0, "2026-05-20 14:30", None),
    (1, None, None),
])
def test_template_signature_respects_preference_and_recorded_date(tmp_path, enabled, approval_date, expected):
    snr = {"posto": "OF-4", "sobrenome": "Silva", "nome": "Ana", "validacao_digital_leave": enabled}
    vacation = {"decidido_em": approval_date}

    def inspect_document(source, destination):
        document = Document(source)
        paragraphs = list(document.paragraphs)
        for table in document.tables:
            paragraphs.extend(report._paragraphs_in_table(table))
        signature = next(p for p in paragraphs if p.text.startswith("12. COORDINATION 2:"))
        drawings = signature._p.xpath(".//w:drawing")
        if enabled:
            assert "Digitally validated by OF-4 SILVA, Ana" in signature.text
            if expected:
                assert f" on {expected}" in signature.text
            else:
                assert "approval date not recorded" in signature.text
            assert len(drawings) == 1
            image_id = drawings[0].xpath(".//a:blip")[0].get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
            assert document.part.related_parts[image_id].blob == (Path(report.DOCS_DIR) / "prt.png").read_bytes()
        else:
            assert "Digitally validated" not in signature.text
            assert not drawings
            assert "____" in signature.text

    with patch.object(report, "_convert_to_pdf", side_effect=inspect_document):
        report.generate_leave_application_pdf(tmp_path / "leave.pdf", {}, vacation, snr)
